"""Copy slides (including pictures) from one PPTX into another — as-is."""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import BinaryIO, Union

from pptx import Presentation
from pptx.oxml.ns import qn

MasterSource = Union[bytes, bytearray, BinaryIO, str, Path]

_R_EMBED = qn("r:embed")
_R_LINK = qn("r:link")


def _as_stream(src: MasterSource) -> BytesIO:
    if isinstance(src, (bytes, bytearray)):
        return BytesIO(src)
    if isinstance(src, (str, Path)):
        return BytesIO(Path(src).read_bytes())
    data = src.read()
    return BytesIO(data)


def _blank_layout(prs: Presentation):
    layouts = prs.slide_layouts
    for layout in layouts:
        name = (layout.name or "").lower()
        if "blank" in name:
            return layout
    return layouts[min(6, len(layouts) - 1)]


def _move_slide(prs: Presentation, from_index: int, to_index: int) -> None:
    sld_id_lst = prs.slides._sldIdLst  # noqa: SLF001
    entries = list(sld_id_lst)
    if from_index < 0 or from_index >= len(entries):
        return
    entry = entries[from_index]
    sld_id_lst.remove(entry)
    to_index = max(0, min(to_index, len(list(sld_id_lst))))
    sld_id_lst.insert(to_index, entry)


def _is_image_rel(rel) -> bool:
    try:
        return "image" in (rel.reltype or "")
    except Exception:
        return False


def _build_image_rid_map(src_slide, dest_slide) -> dict[str, str]:
    """Map source slide image rIds → newly related rIds on dest slide."""
    rid_map: dict[str, str] = {}
    for rel in src_slide.part.rels.values():
        if not _is_image_rel(rel):
            continue
        try:
            blob = rel.target_part.blob
        except Exception:
            continue
        try:
            _part, new_rid = dest_slide.part.get_or_add_image_part(BytesIO(blob))
        except Exception:
            continue
        rid_map[rel.rId] = new_rid
    return rid_map


def _rebind_rids(element, rid_map: dict[str, str]) -> None:
    if not rid_map:
        return
    for el in element.iter():
        for attr in (_R_EMBED, _R_LINK):
            old = el.get(attr)
            if old and old in rid_map:
                el.set(attr, rid_map[old])


def _copy_slide_background(src_slide, dest_slide, rid_map: dict[str, str]) -> None:
    """Copy cSld background (solid / gradient / blip) when present."""
    try:
        src_bg = src_slide._element.cSld.find(qn("p:bg"))  # noqa: SLF001
    except Exception:
        src_bg = None
    if src_bg is None:
        return
    dest_cSld = dest_slide._element.cSld  # noqa: SLF001
    existing = dest_cSld.find(qn("p:bg"))
    if existing is not None:
        dest_cSld.remove(existing)
    new_bg = deepcopy(src_bg)
    _rebind_rids(new_bg, rid_map)
    # Background must sit before spTree
    sp_tree = dest_cSld.find(qn("p:spTree"))
    if sp_tree is not None:
        sp_tree.addprevious(new_bg)
    else:
        dest_cSld.append(new_bg)


def _clear_slide_shapes(slide) -> None:
    for shape in list(slide.shapes):
        try:
            shape._element.getparent().remove(shape._element)  # noqa: SLF001
        except Exception:
            pass


def _copy_one_slide(dest_prs: Presentation, src_slide) -> int:
    """Append a visual copy of src_slide onto dest_prs. Returns new slide index."""
    dest_slide = dest_prs.slides.add_slide(_blank_layout(dest_prs))
    _clear_slide_shapes(dest_slide)

    rid_map = _build_image_rid_map(src_slide, dest_slide)
    _copy_slide_background(src_slide, dest_slide, rid_map)

    for shape in src_slide.shapes:
        try:
            new_el = deepcopy(shape.element)
        except Exception:
            continue
        _rebind_rids(new_el, rid_map)
        dest_slide.shapes._spTree.insert_element_before(new_el, "p:extLst")  # noqa: SLF001

    return len(dest_prs.slides) - 1


def insert_pptx_after(
    dest_prs: Presentation,
    after_index: int,
    source: MasterSource,
) -> int:
    """
    Insert every slide from `source` immediately after `after_index` in dest_prs.
    Slides are copied as-is (shapes + pictures + slide background).
    Returns number of slides inserted.
    """
    src_prs = Presentation(_as_stream(source))
    if not src_prs.slides:
        return 0

    n = 0
    for src_slide in src_prs.slides:
        _copy_one_slide(dest_prs, src_slide)
        n += 1

    # Newly added slides sit at the end — move them after `after_index` in order.
    start = len(dest_prs.slides) - n
    for i in range(n):
        _move_slide(dest_prs, start + i, after_index + 1 + i)
    return n


def find_named_token_slide(prs: Presentation, token: str) -> int | None:
    """Return first slide index whose shape name or {{TOKEN}} text matches."""
    bare = "{{" + token + "}}"
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            name = (getattr(shape, "name", None) or "").strip()
            if name == token:
                return i
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text or ""
            if bare in text or text.strip() == token:
                return i
    return None
