"""Second-pass PPT layout: reorder, recanter, and fit text like the HTML deck.

After slides are painted, each text box is named ``role:*``. This pass
snaps them to the reading column, stacks them with even gaps, optically
centers short clusters (cover / section / lyrics), and shrinks type so
nothing overflows the rose band at the bottom.
"""

from __future__ import annotations

import re

from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

from build_master_templates import CONTENT_W, MARGIN_L, SLIDE_H

_POLISH_VERSION = "2026-09-04-pptx-refine-v1"

_EMU = 914400.0
_TOKEN_RE = re.compile(r"\{\{[A-Za-z0-9_]+\}\}")

# Stay above the corner roses (y ≈ 6.60")
_Y_TOP = 0.34
_Y_BOTTOM = 6.36
_GAP = 0.12
_HEADER_H = 0.46


def _inch(emu) -> float:
    try:
        return float(emu) / _EMU
    except Exception:
        return 0.0


def _role(shape) -> str:
    name = (getattr(shape, "name", None) or "")
    if name.startswith("role:"):
        return name.split(":", 1)[1].strip()
    return ""


def _roles(slide) -> dict:
    found: dict = {}
    for shape in slide.shapes:
        role = _role(shape)
        if role:
            found[role] = shape
    return found


def _all_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    parts = []
    for para in shape.text_frame.paragraphs:
        parts.append("".join(run.text or "" for run in para.runs) or (para.text or ""))
    return "\n".join(parts).strip()


def _visual_len(text: str) -> float:
    n = 0.0
    for ch in text or "":
        n += 1.0 if ord(ch) > 127 else 0.55
    return n


def _place(shape, left: float, top: float, width: float, height: float) -> None:
    shape.left = Inches(max(0.04, left))
    shape.top = Inches(max(0.04, top))
    shape.width = Inches(max(0.04, width))
    shape.height = Inches(max(0.012, height))


def _set_anchor(shape, anchor) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    try:
        shape.text_frame.word_wrap = True
        shape.text_frame.vertical_anchor = anchor
    except Exception:
        pass


def _current_pt(shape) -> float:
    if not getattr(shape, "has_text_frame", False):
        return 28.0
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            sz = getattr(run.font, "size", None)
            if sz is not None:
                try:
                    return float(sz.pt)
                except Exception:
                    pass
    return 28.0


def _set_pt(shape, size_pt: float) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    pt = Pt(max(14.0, size_pt))
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            try:
                run.font.size = pt
            except Exception:
                pass


def _scrub_tokens(shape) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.text and _TOKEN_RE.search(run.text):
                run.text = _TOKEN_RE.sub("", run.text).strip()


def _est_height(text: str, width_in: float, font_pt: float, *, line_spacing: float = 1.38) -> float:
    usable = max(1.0, width_in - 0.16)
    cpl = max(4.0, usable * 72.0 / max(12.0, font_pt))
    lines = 0
    raw_lines = (text or "").split("\n") or [""]
    for raw in raw_lines:
        vis = _visual_len(raw.strip()) if raw.strip() else 0.5
        lines += max(1, int((vis + cpl - 0.01) // cpl))
    h_pt = lines * font_pt * line_spacing + max(0, len(raw_lines) - 1) * 3.0
    return h_pt / 72.0 + 0.10


def _fit_pt(shape, width_in: float, max_h: float, *, min_pt: float, max_pt: float | None = None) -> float:
    text = _all_text(shape)
    start = _current_pt(shape)
    hi = max_pt if max_pt is not None else start
    hi = max(min_pt, min(hi, start) if max_pt is None else hi)
    if not text:
        return start
    chosen = min_pt
    for pt in range(int(round(hi)), int(round(min_pt)) - 1, -1):
        if _est_height(text, width_in, float(pt)) <= max_h:
            chosen = float(pt)
            break
    _set_pt(shape, chosen)
    return chosen


def _needed(shape, width_in: float, font_pt: float | None = None) -> float:
    text = _all_text(shape)
    pt = font_pt if font_pt is not None else _current_pt(shape)
    return max(0.28, _est_height(text, width_in, pt))


def _bring_roles_front(slide) -> None:
    sp_tree = slide.shapes._spTree  # noqa: SLF001
    tagged = [shape._element for shape in slide.shapes if _role(shape)]  # noqa: SLF001
    for el in tagged:
        sp_tree.remove(el)
    for el in tagged:
        sp_tree.append(el)


def _column_box(shape, top: float, height: float, *, left: float | None = None, width: float | None = None) -> None:
    _place(shape, left if left is not None else _inch(MARGIN_L), top, width if width is not None else _inch(CONTENT_W), height)


def _center_stack(items: list[tuple], y_lo: float, y_hi: float, *, gap: float = _GAP) -> None:
    """items: (shape, height_in) — skip missing/empty shapes."""
    live = [(s, h) for s, h in items if s is not None and h > 0]
    if not live:
        return
    total = sum(h for _, h in live) + gap * max(0, len(live) - 1)
    room = max(0.3, y_hi - y_lo)
    top = y_lo + max(0.0, (room - total) / 2.0)
    col_l = _inch(MARGIN_L)
    col_w = _inch(CONTENT_W)
    y = top
    for shape, height in live:
        _column_box(shape, y, height, left=col_l, width=col_w)
        y += height + gap


def _layout_cover(roles: dict) -> None:
    footer = roles.get("footer")
    foot_h = 0.42 if footer else 0.0
    y_hi = _Y_BOTTOM - (0.18 if footer else 0.0)
    cluster = []
    accent = roles.get("accent")
    if accent is not None:
        cluster.append((accent, 0.10))
    for key, default_h in (("title", 1.15), ("subtitle", 0.62), ("extra", 0.52)):
        shape = roles.get(key)
        if shape is None or not _all_text(shape):
            continue
        width = _inch(CONTENT_W)
        pt = _fit_pt(shape, width, default_h + 0.35, min_pt=22, max_pt=_current_pt(shape))
        cluster.append((shape, min(default_h + 0.25, _needed(shape, width, pt))))
        _set_anchor(shape, MSO_ANCHOR.MIDDLE)
    _center_stack(cluster, _Y_TOP + 0.35, y_hi, gap=0.16)
    if accent is not None:
        # Keep the gold tick short and centered in its row
        _place(accent, _inch(MARGIN_L) + (_inch(CONTENT_W) - 1.55) / 2.0, _inch(accent.top), 1.55, 0.055)
    if footer is not None:
        _column_box(footer, _inch(SLIDE_H) - 0.58, foot_h)
        _set_anchor(footer, MSO_ANCHOR.MIDDLE)


def _layout_section(roles: dict) -> None:
    cluster = []
    for key, default_h, min_pt in (("title", 1.25, 28), ("subtitle", 0.95, 22)):
        shape = roles.get(key)
        if shape is None or not _all_text(shape):
            continue
        width = _inch(CONTENT_W)
        pt = _fit_pt(shape, width, default_h + 0.4, min_pt=min_pt, max_pt=_current_pt(shape))
        cluster.append((shape, min(default_h + 0.35, _needed(shape, width, pt) + 0.08)))
        _set_anchor(shape, MSO_ANCHOR.MIDDLE)
    _center_stack(cluster, _Y_TOP + 0.55, _Y_BOTTOM - 0.15, gap=0.22)


def _header_row(roles: dict) -> float:
    left = roles.get("header")
    right = roles.get("header_right")
    col_l = _inch(MARGIN_L)
    col_w = _inch(CONTENT_W)
    y = _Y_TOP
    if left is None and right is None:
        return y
    if left is not None and right is not None:
        half = col_w * 0.50
        _fit_pt(left, half - 0.08, _HEADER_H, min_pt=14, max_pt=18)
        _fit_pt(right, half - 0.08, _HEADER_H, min_pt=14, max_pt=18)
        _place(left, col_l, y, half - 0.08, _HEADER_H)
        _place(right, col_l + half + 0.08, y, half - 0.08, _HEADER_H)
        _set_anchor(left, MSO_ANCHOR.MIDDLE)
        _set_anchor(right, MSO_ANCHOR.MIDDLE)
    elif left is not None:
        _column_box(left, y, _HEADER_H)
        _set_anchor(left, MSO_ANCHOR.MIDDLE)
    else:
        _column_box(right, y, _HEADER_H)
        _set_anchor(right, MSO_ANCHOR.MIDDLE)
    rule = roles.get("rule")
    if rule is not None:
        _place(rule, col_l, y + _HEADER_H + 0.02, col_w, 0.018)
        return y + _HEADER_H + 0.12
    return y + _HEADER_H + 0.10


def _layout_body_in_band(
    body,
    y_lo: float,
    *,
    min_pt: float,
    max_pt: float,
    center: bool,
) -> None:
    if body is None:
        return
    width = _inch(CONTENT_W)
    band = max(0.8, _Y_BOTTOM - y_lo)
    pt = _fit_pt(body, width, band, min_pt=min_pt, max_pt=max_pt)
    need = min(band, _needed(body, width, pt) + 0.06)
    if center and need < band * 0.92:
        top = y_lo + (band - need) / 2.0
        _column_box(body, top, need)
        _set_anchor(body, MSO_ANCHOR.MIDDLE)
    else:
        _column_box(body, y_lo, band)
        _set_anchor(body, MSO_ANCHOR.TOP)


def _layout_lyric(roles: dict) -> None:
    y = _header_row(roles)
    _layout_body_in_band(roles.get("body"), y + 0.06, min_pt=26, max_pt=40, center=True)


def _layout_creed(roles: dict) -> None:
    y = _header_row(roles)
    title = roles.get("title")
    body = roles.get("body")
    width = _inch(CONTENT_W)
    band_lo = y + 0.08
    items = []
    if title is not None and _all_text(title):
        pt = _fit_pt(title, width, 1.15, min_pt=26, max_pt=_current_pt(title))
        items.append((title, min(1.15, _needed(title, width, pt) + 0.06)))
        _set_anchor(title, MSO_ANCHOR.MIDDLE)
    if body is not None and _all_text(body):
        # Body shares remaining band after title in the centered stack
        remain = max(1.4, _Y_BOTTOM - band_lo - (items[0][1] + _GAP if items else 0))
        pt = _fit_pt(body, width, remain, min_pt=24, max_pt=min(34, _current_pt(body)))
        items.append((body, min(remain, _needed(body, width, pt) + 0.08)))
        _set_anchor(body, MSO_ANCHOR.MIDDLE)
    _center_stack(items, band_lo, _Y_BOTTOM, gap=0.18)


def _layout_scripture(roles: dict, *, center: bool = False) -> None:
    y = _header_row(roles)
    title = roles.get("title")
    if title is not None and _all_text(title):
        _column_box(title, y, 0.62)
        _set_anchor(title, MSO_ANCHOR.MIDDLE)
        y += 0.70
    _layout_body_in_band(
        roles.get("body"),
        y + 0.04,
        min_pt=22,
        max_pt=34,
        center=center,
    )


def _layout_sermon(roles: dict) -> None:
    y = _header_row(roles)
    footer = roles.get("footer")
    y_hi = _Y_BOTTOM - (0.12 if footer else 0.0)
    cluster = []
    kicker = roles.get("kicker")
    if kicker is not None and _all_text(kicker):
        cluster.append((kicker, 0.42))
        _set_anchor(kicker, MSO_ANCHOR.MIDDLE)
    title = roles.get("title")
    if title is not None and _all_text(title):
        width = _inch(CONTENT_W)
        pt = _fit_pt(title, width, 2.2, min_pt=28, max_pt=_current_pt(title))
        cluster.append((title, min(2.15, _needed(title, width, pt) + 0.12)))
        _set_anchor(title, MSO_ANCHOR.MIDDLE)
    sub = roles.get("subtitle")
    if sub is not None and _all_text(sub):
        cluster.append((sub, 0.50))
        _set_anchor(sub, MSO_ANCHOR.MIDDLE)
    _center_stack(cluster, y + 0.15, y_hi, gap=0.16)
    if footer is not None:
        _column_box(footer, _inch(SLIDE_H) - 0.58, 0.40)
        _set_anchor(footer, MSO_ANCHOR.MIDDLE)


def _layout_list(roles: dict) -> None:
    y = _Y_TOP + 0.08
    title = roles.get("title")
    sub = roles.get("subtitle")
    body = roles.get("body")
    if title is not None:
        _column_box(title, y, 0.72)
        _set_anchor(title, MSO_ANCHOR.MIDDLE)
        y += 0.78
    if sub is not None and _all_text(sub):
        _column_box(sub, y, 0.42)
        _set_anchor(sub, MSO_ANCHOR.MIDDLE)
        y += 0.50
    _layout_body_in_band(body, y + 0.04, min_pt=16, max_pt=26, center=False)


def _layout_centered_title(roles: dict) -> None:
    cluster = []
    for key, default_h in (("title", 1.2), ("subtitle", 0.9)):
        shape = roles.get(key)
        if shape is None or not _all_text(shape):
            continue
        width = _inch(CONTENT_W)
        pt = _fit_pt(shape, width, default_h + 0.4, min_pt=22, max_pt=_current_pt(shape))
        cluster.append((shape, min(default_h + 0.3, _needed(shape, width, pt) + 0.08)))
        _set_anchor(shape, MSO_ANCHOR.MIDDLE)
    _center_stack(cluster, _Y_TOP + 0.55, _Y_BOTTOM - 0.1, gap=0.2)


def refine_slide(slide, spec: dict | None = None) -> None:
    """Reorder and refit one slide after it has been painted."""
    spec = spec or {}
    kind = (spec.get("type") or "title").lower()
    title = (spec.get("title") or "").strip()
    content = spec.get("content") or ""
    roles = _roles(slide)

    for shape in list(roles.values()):
        _scrub_tokens(shape)

    if kind == "cover":
        _layout_cover(roles)
    elif kind == "section":
        _layout_section(roles)
    elif kind == "lyric":
        _layout_lyric(roles)
    elif kind == "creed":
        _layout_creed(roles)
    elif kind == "scripture":
        _layout_scripture(roles, center=False)
    elif kind == "responsive":
        _layout_scripture(roles, center=True)
    elif kind == "sermon":
        _layout_sermon(roles)
    elif title == "예배 순서" or "order-list" in content or title.startswith("11.") or title.startswith("12."):
        _layout_list(roles)
    else:
        _layout_centered_title(roles)

    _bring_roles_front(slide)


def refine_presentation(prs, specs: list[dict] | None = None) -> None:
    slides = list(prs.slides)
    specs = specs or [{}] * len(slides)
    for slide, spec in zip(slides, specs):
        try:
            refine_slide(slide, spec)
        except Exception:
            continue
